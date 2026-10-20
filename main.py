import os
import sys
import random
import msvcrt
from datetime import datetime
from openpyxl import Workbook, load_workbook

JOURNAL_FILE = "journal.xlsx"
HEADERS = ["Last name", "First name", "Номер билета", "Дата и время"]


def read_line_with_esc(prompt: str) -> str | None:
    """
    Посимвольное чтение строки из консоли с поддержкой клавиши ESC.
    Возвращает введённую строку или None, если нажата клавиша ESC.
    """
    sys.stdout.write(prompt)
    sys.stdout.flush()
    chars: list[str] = []

    while True:
        ch = msvcrt.getwch()

        # ESC — выход из приложения
        if ch == "\x1b":
            sys.stdout.write("\n")
            sys.stdout.flush()
            return None

        # Ctrl+C — генерация KeyboardInterrupt
        if ch == "\x03":
            sys.stdout.write("\n")
            sys.stdout.flush()
            raise KeyboardInterrupt

        # Enter — завершение ввода строки
        if ch in ("\r", "\n"):
            sys.stdout.write("\n")
            sys.stdout.flush()
            return "".join(chars)

        # Backspace — удаление последнего символа
        if ch == "\x08":
            if chars:
                chars.pop()
                sys.stdout.write("\b \b")
                sys.stdout.flush()
            continue

        # Специальные клавиши (стрелки, F1-F12 и др.) начинаются с \x00 или \xe0
        if ch in ("\x00", "\xe0"):
            msvcrt.getwch()  # Пропускаем второй код клавиши
            continue

        # Обычный печатный символ
        if ch.isprintable():
            chars.append(ch)
            sys.stdout.write(ch)
            sys.stdout.flush()


def prompt_non_empty(label: str) -> str | None:
    """
    Запрашивает непустую строку (пробелы по краям обрезаются).
    Если нажата ESC, возвращает None.
    """
    while True:
        raw = read_line_with_esc(label)
        if raw is None:
            return None
        value = raw.strip()
        if value:
            return value
        print("Ошибка: поле не может быть пустым. Повторите ввод.")


def append_to_excel(last_name: str, first_name: str, ticket_num: int, timestamp: str) -> bool:
    """
    Дописывает запись в файл journal.xlsx и сразу сохраняет его.
    Если файл заблокирован (например, открыт в Excel), предлагает закрыть его и повторить.
    Возвращает True при успешном сохранении или False при отмене по ESC.
    """
    while True:
        try:
            if os.path.exists(JOURNAL_FILE):
                wb = load_workbook(JOURNAL_FILE)
                ws = wb.active
            else:
                wb = Workbook()
                ws = wb.active
                ws.title = "Journal"
                ws.append(HEADERS)

            ws.append([last_name, first_name, ticket_num, timestamp])
            wb.save(JOURNAL_FILE)
            wb.close()
            return True
        except PermissionError:
            print(
                f"\n[Ошибка] Не удалось записать данные: файл '{JOURNAL_FILE}' открыт в другой программе (например, в Excel)."
            )
            res = read_line_with_esc(
                "Пожалуйста, закройте файл в Excel и нажмите Enter для повторной попытки (или ESC для выхода): "
            )
            if res is None:
                return False
        except Exception as e:
            print(f"\n[Ошибка] Ошибка при работе с файлом '{JOURNAL_FILE}': {e}")
            res = read_line_with_esc(
                "Нажмите Enter для повторной попытки (или ESC для выхода): "
            )
            if res is None:
                return False


def main() -> None:
    print("Для выхода нажмите ESC.")
    try:
        while True:
            last_name = prompt_non_empty("Last name: ")
            if last_name is None:
                print("Работа приложения завершена.")
                break

            first_name = prompt_non_empty("First name: ")
            if first_name is None:
                print("Работа приложения завершена.")
                break

            ticket_num = random.randint(1, 20)
            print(f"Билет № {ticket_num}")

            timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            saved = append_to_excel(last_name, first_name, ticket_num, timestamp)
            if not saved:
                print("Работа приложения завершена.")
                break
    except KeyboardInterrupt:
        print("\nРабота приложения прервана пользователем.")


if __name__ == "__main__":
    main()
