from pathlib import Path
from docx import Document
from openpyxl import Workbook, load_workbook
import pytest

from ticket_service import (
    Ticket,
    append_result_row,
    assign_ticket,
    load_students,
    parse_tickets,
)


def test_parse_tickets_valid_and_malformed(tmp_path: Path) -> None:
    """
    Тест 1: Парсинг билетов из Word (tickets.docx).
    Проверяет:
    - Корректное чтение валидных билетов с 3 вопросами.
    - Пропуск билета, у которого меньше 3 вопросов (без падения приложения).
    - Пропуск билета с некорректным заголовком.
    """
    docx_path = tmp_path / "tickets.docx"
    doc = Document()

    # Валидный Билет 1
    doc.add_paragraph("Билет 1")
    doc.add_paragraph("1. Вопрос первый")
    doc.add_paragraph("2. Вопрос второй")
    doc.add_paragraph("3. Вопрос третий")

    # Невалидный Билет 2 (всего 2 вопроса — должен быть пропущен)
    doc.add_paragraph("Билет 2")
    doc.add_paragraph("1. Только первый вопрос")
    doc.add_paragraph("2. Только второй вопрос")

    # Некорректный заголовок (должен быть пропущен)
    doc.add_paragraph("Билет X_некорректный")
    doc.add_paragraph("1. Вопрос А")
    doc.add_paragraph("2. Вопрос Б")
    doc.add_paragraph("3. Вопрос В")

    # Валидный Билет 3
    doc.add_paragraph("Билет 3")
    doc.add_paragraph("1. Третий билет вопрос 1")
    doc.add_paragraph("2. Третий билет вопрос 2")
    doc.add_paragraph("3. Третий билет вопрос 3")

    doc.save(str(docx_path))

    tickets = parse_tickets(str(docx_path))

    assert len(tickets) == 2
    assert 1 in tickets
    assert 2 not in tickets
    assert 3 in tickets
    assert tickets[1].questions == [
        "1. Вопрос первый",
        "2. Вопрос второй",
        "3. Вопрос третий",
    ]


def test_load_students_by_group(tmp_path: Path) -> None:
    """
    Тест 2: Поиск студентов по группе из Excel (students.xlsx), включая пустую группу.
    """
    excel_path = tmp_path / "students.xlsx"
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Группа-А"
    ws1.append(["Фамилия", "Имя"])
    ws1.append(["Иванов", "Иван"])
    ws1.append(["Петрова", "Мария"])

    ws2 = wb.create_sheet("Пустая-Группа")
    ws2.append(["Фамилия", "Имя"])

    wb.save(str(excel_path))
    wb.close()

    groups = load_students(str(excel_path))

    assert "Группа-А" in groups
    assert "Пустая-Группа" in groups
    assert len(groups["Группа-А"]) == 2
    assert groups["Группа-А"][0].full_name == "Иванов Иван"
    assert groups["Группа-А"][1].full_name == "Петрова Мария"
    assert groups["Пустая-Группа"] == []


def test_idempotency_same_student_gets_same_ticket(tmp_path: Path) -> None:
    """
    Тест 3: Идемпотентность (повторный выбор того же студента возвращает тот же билет,
    в журнал дописывается новая строка с тем же номером и пометкой 'да' в колонке 'Повтор').
    """
    results_path = tmp_path / "results.xlsx"
    tickets = {
        i: Ticket(number=i, questions=[f"1. Q1_{i}", f"2. Q2_{i}", f"3. Q3_{i}"])
        for i in range(1, 10)
    }

    # Первая генерация для студента
    first_ticket, is_repeat_1 = assign_ticket(
        group="ИВТ-101",
        last_name="Иванов",
        first_name="Иван",
        tickets=tickets,
        results_filepath=str(results_path),
    )
    assert is_repeat_1 is False
    append_result_row(
        group="ИВТ-101",
        last_name="Иванов",
        first_name="Иван",
        ticket_number=first_ticket.number,
        is_repeat=is_repeat_1,
        results_filepath=str(results_path),
    )

    # Повторный выбор того же студента
    second_ticket, is_repeat_2 = assign_ticket(
        group="ИВТ-101",
        last_name="Иванов",
        first_name="Иван",
        tickets=tickets,
        results_filepath=str(results_path),
    )
    assert is_repeat_2 is True
    assert second_ticket.number == first_ticket.number
    append_result_row(
        group="ИВТ-101",
        last_name="Иванов",
        first_name="Иван",
        ticket_number=second_ticket.number,
        is_repeat=is_repeat_2,
        results_filepath=str(results_path),
    )

    # Проверяем содержимое results.xlsx
    wb = load_workbook(str(results_path))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    # Шапка + 2 записи
    assert len(rows) == 3
    assert rows[1][3] == first_ticket.number
    assert rows[1][5] == "нет"
    assert rows[2][3] == first_ticket.number
    assert rows[2][5] == "да"
